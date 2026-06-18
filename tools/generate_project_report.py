from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "综合项目报告-CampusHub.docx"


def set_east_asian_font(run, name="宋体"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")


def set_para_format(paragraph, first_line=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    paragraph.alignment = align
    pf = paragraph.paragraph_format
    pf.line_spacing = 1.0
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    if first_line:
        pf.first_line_indent = Pt(21)


def add_run(paragraph, text, size=10.5, bold=False, font="宋体", color=None):
    run = paragraph.add_run(text)
    set_east_asian_font(run, font)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def add_paragraph(doc, text=""):
    p = doc.add_paragraph()
    set_para_format(p, first_line=True)
    add_run(p, text, size=10.5)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.line_spacing = 1.0
    pf.space_before = Pt(6 if level == 1 else 3)
    pf.space_after = Pt(6 if level == 1 else 3)
    pf.first_line_indent = Pt(0)
    size = 16 if level == 1 else 14 if level == 2 else 12
    bold = True
    run = add_run(p, text, size=size, bold=bold, font="宋体")
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.line_spacing = 1.0
    pf.space_before = Pt(3)
    pf.space_after = Pt(6)
    add_run(p, text, size=9, font="宋体")
    return p


def set_cell_text(cell, text, bold=False):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(0)
    p.text = ""
    add_run(p, text, size=12, bold=bold)


def set_table_borders(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "8")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), "000000")
        borders.append(tag)
    tbl_pr.append(borders)


def add_info_table(doc):
    table = doc.add_table(rows=6, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    labels = ["项目名称", "学    号", "姓    名", "班    级", "学    院", "完成日期"]
    values = ["CampusHub 校园综合社区平台", "", "", "", "信息工程学院", "2026 年 6 月"]
    for i, (label, value) in enumerate(zip(labels, values)):
        set_cell_text(table.cell(i, 0), label, bold=True)
        set_cell_text(table.cell(i, 1), value)
        table.cell(i, 0).width = Cm(4)
        table.cell(i, 1).width = Cm(8)
    return table


def add_simple_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    if widths is None:
        widths = [Cm(16 / len(headers))] * len(headers)
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.width = widths[i]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.text = ""
        add_run(p, h, size=10.5, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].width = widths[i]
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if len(str(value)) > 12 else WD_ALIGN_PARAGRAPH.CENTER
            p.text = ""
            add_run(p, str(value), size=10.5)
    return table


def add_cover(doc):
    for text, size, bold, before in [
        ("江西现代职业技术学院", 18, True, 2),
        ("《Java Web应用开发》", 18, True, 1),
        ("综合项目报告", 22, True, 1),
    ]:
        for _ in range(before):
            doc.add_paragraph()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        add_run(p, text, size=size, bold=bold, font="宋体")
    for _ in range(3):
        doc.add_paragraph()
    add_info_table(doc)
    doc.add_page_break()


def add_list(doc, items):
    for item in items:
        p = doc.add_paragraph()
        p.style = doc.styles["Normal"]
        p.paragraph_format.left_indent = Pt(21)
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(0)
        add_run(p, "（1）" if False else "", size=10.5)
        add_run(p, item, size=10.5)


def build_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)

    add_cover(doc)

    add_heading(doc, "第一章 项目简介", 1)
    add_heading(doc, "1.1 项目背景", 2)
    add_paragraph(doc, "随着校园信息化建设不断推进，学生在校内学习、生活和社交过程中会产生大量分散需求，例如发布校园动态、转让闲置物品、寻找遗失物品、报名校园活动、接收通知公告以及与同学私信沟通等。传统做法通常依赖多个独立渠道，信息分散、查询不便，管理员也难以及时进行内容维护和违规处理。CampusHub 校园综合社区平台围绕学生日常高频场景，将校园广场、二手市场、失物招领、活动报名、站内消息、全站搜索和后台管理整合到同一 Java Web 应用中，提升信息流转效率和系统管理能力。")
    add_heading(doc, "1.2 项目目标", 2)
    add_paragraph(doc, "本项目目标是完成一个基于 JSP、Servlet、JDBC 和 MySQL 的校园综合社区系统。系统面向学生端和管理员端两个角色：学生端强调信息发布、互动、交易和个人中心；管理员端强调数据统计、用户管理、内容审核和举报处理。项目实现过程中坚持传统 Java Web 分层思想，将请求处理、业务规则、数据访问和页面展示分离，便于维护、扩展和课程答辩展示。")
    add_heading(doc, "1.3 功能概述", 2)
    add_simple_table(doc, ["功能模块", "主要功能"], [
        ["用户与认证", "注册、验证码登录、退出登录、记住我、权限过滤、账号注销"],
        ["校园广场", "帖子发布、列表浏览、详情查看、评论、点赞、收藏、举报"],
        ["二手市场", "商品发布、分类筛选、编辑、收藏、状态管理、卖家私信、模拟支付"],
        ["失物招领", "失物/招领发布、详情查看、认领申请、认领审核和状态流转"],
        ["校园活动", "活动发布、报名、取消报名、人数限制、截止时间校验"],
        ["消息与搜索", "站内通知、未读统计、私信会话、跨模块全站搜索"],
        ["后台管理", "统计面板、用户管理、内容管理、公告维护、举报处理"],
    ], [Cm(4), Cm(12)])
    add_caption(doc, "表1.1 CampusHub 功能模块概览")
    add_heading(doc, "1.4 开发环境与技术选型", 2)
    add_simple_table(doc, ["类别", "技术"], [
        ["后端语言", "Java 21"],
        ["Web 技术", "Servlet 4.0、JSP、Filter"],
        ["数据访问", "JDBC、PreparedStatement、事务控制"],
        ["数据库", "MySQL 8"],
        ["前端", "HTML、CSS、JavaScript、Fetch API"],
        ["构建工具", "Maven"],
        ["服务器", "Apache Tomcat 9"],
        ["安全相关", "BCrypt、Session、Cookie、验证码、Remember Me Token"],
        ["测试", "JUnit 5"],
    ], [Cm(4), Cm(12)])
    add_caption(doc, "表1.2 项目开发环境与技术选型")
    add_caption(doc, "图1.1 系统首页运行效果图（截图待补）")
    add_caption(doc, "图1.2 登录与注册页面运行效果图（截图待补）")

    add_heading(doc, "第二章 架构设计", 1)
    add_heading(doc, "2.1 总体架构", 2)
    add_paragraph(doc, "CampusHub 采用典型的 Java Web 分层架构，整体请求链路为：浏览器发起请求，Filter 统一处理编码、登录恢复和权限校验，Servlet 接收并分发请求，Service 负责业务规则和事务组织，DAO 使用 JDBC 与 MySQL 交互，最终由 JSP 页面或 JSON 响应返回结果。该架构没有引入 Spring、Vue 或 React，符合 Java Web 课程对 Servlet、JSP、JDBC 基础能力训练的要求。")
    add_simple_table(doc, ["层次", "职责", "代表文件"], [
        ["视图层", "负责页面展示、表单提交、前端交互和局部刷新", "JSP、CSS、JavaScript"],
        ["控制层", "接收 HTTP 请求、解析参数、调用业务服务、转发视图或输出 JSON", "HomeServlet、GoodsCreateServlet、AdminServlet"],
        ["业务层", "进行参数校验、状态判断、业务流程编排和异常处理", "UserService、GoodsService、ActivityService"],
        ["数据访问层", "封装 SQL、映射结果集、执行事务和持久化操作", "JdbcUserDao、JdbcPostDao、JdbcAdminDao"],
        ["模型层", "承载实体、视图对象和操作结果", "User、Post、Goods、ServiceResult"],
    ], [Cm(2.8), Cm(7), Cm(6.2)])
    add_caption(doc, "表2.1 系统分层结构说明")
    add_heading(doc, "2.2 请求处理流程", 2)
    add_paragraph(doc, "以发布帖子为例，用户在页面填写标题、正文、分类和图片后，前端通过表单或异步请求提交到对应 Servlet；AuthFilter 首先确认用户已登录，Servlet 获取 SessionUser 和请求参数后调用 PostService；PostService 负责校验标题长度、正文内容、分类有效性和图片字段；校验通过后调用 JdbcPostDao 写入 posts 表；最后 Servlet 根据执行结果返回页面跳转或 JSON 提示。类似流程也应用于商品、失物招领、活动、举报、私信等模块。")
    add_caption(doc, "图2.1 系统总体架构图（截图或绘图待补）")
    add_caption(doc, "图2.2 用户请求处理流程图（截图或绘图待补）")
    add_heading(doc, "2.3 权限与安全设计", 2)
    add_paragraph(doc, "系统通过 EncodingFilter 统一请求和响应编码，避免中文参数乱码；RememberMeFilter 在会话缺失时尝试根据持久化令牌恢复登录状态；AuthFilter 保护个人中心、收藏、私信和发布等需要登录的路径；AdminAuthFilter 保护后台管理路径。密码使用 BCrypt 哈希保存，数据库访问统一采用 PreparedStatement，减少 SQL 注入风险。配置方面，数据库密码不再提交到仓库，运行时优先从环境变量或 JVM 参数读取。")
    add_heading(doc, "2.4 项目目录结构", 2)
    add_paragraph(doc, "项目采用 Maven 标准目录结构，Java 源码集中在 src/main/java/cn/campushub 下，按 config、constant、dao、filter、model、service、servlet、util 进行模块划分；JSP、CSS、JavaScript 和图片资源位于 src/main/webapp；数据库初始化脚本位于 src/main/resources/database/schema.sql；单元测试位于 src/test/java。当前项目约包含 181 个 Java 文件、28 个 JSP 文件、14 个 CSS 文件、14 个 JavaScript 文件，结构清晰，便于定位和维护。")

    add_heading(doc, "第三章 数据访问层设计", 1)
    add_heading(doc, "3.1 数据库设计概述", 2)
    add_paragraph(doc, "数据库名称为 campushub，完整初始化脚本为 src/main/resources/database/schema.sql。当前数据库包含 users、categories、posts、comments、likes、favorites、goods、goods_orders、lost_found、claim_requests、activities、activity_registrations、notices、checkins、messages、reports、private_conversations、private_messages、remember_tokens、user_experience_logs、account_cancel_logs 等 21 张主要业务表。users 表是中心实体，与帖子、评论、收藏、商品、订单、失物招领、活动报名、消息、举报和私信等表建立关联。")
    add_heading(doc, "3.2 主要数据表设计", 2)
    add_simple_table(doc, ["数据表", "设计说明"], [
        ["users", "存储用户账号、密码哈希、昵称、头像、学籍信息、角色、状态、经验值和注销信息"],
        ["posts/comments/likes/favorites", "支撑校园广场的内容发布、评论、点赞和收藏互动"],
        ["goods/goods_orders", "支撑二手市场商品发布、交易方式选择和模拟二维码支付订单"],
        ["lost_found/claim_requests", "支撑失物招领信息发布、认领申请和处理流程"],
        ["activities/activity_registrations", "支撑活动发布、报名、取消报名和人数限制"],
        ["messages/private_*", "支撑系统通知、未读统计和一对一私信沟通"],
        ["reports/notices", "支撑内容举报、后台审核和校园公告维护"],
    ], [Cm(5), Cm(11)])
    add_caption(doc, "表3.1 主要数据表说明")
    add_caption(doc, "图3.1 数据库关系图（截图待补）")
    add_heading(doc, "3.3 DAO 接口与 JDBC 实现", 2)
    add_paragraph(doc, "数据访问层采用接口加实现类的方式组织。例如 UserDao 定义用户查询、注册和重复性校验能力，JdbcUserDao 使用 JDBC 完成具体 SQL；GoodsDao 与 JdbcGoodsDao 封装商品列表、详情、发布、收藏和状态变更；AdminDao 与 JdbcAdminDao 封装后台统计、用户管理、内容管理、公告管理和举报处理。业务层依赖 DAO 接口，使核心业务逻辑不直接散落 SQL。")
    add_heading(doc, "3.4 事务处理设计", 2)
    add_paragraph(doc, "对于涉及多表一致性的业务，DAO 层使用 Connection#setAutoCommit(false) 显式开启事务，并在成功后提交、异常时回滚。例如模拟支付确认需要同时更新 goods_orders 支付状态和 goods 商品状态；活动报名需要锁定活动记录、写入报名记录、更新 current_members 并发送消息；账号注销需要隐藏用户内容、清理个人互动数据、处理订单状态并写入注销日志。这些事务设计保证了并发和异常情况下的数据一致性。")
    add_heading(doc, "3.5 数据库配置与安全", 2)
    add_paragraph(doc, "DatabaseConfig 统一读取数据库驱动、连接地址、用户名和密码。配置读取顺序为环境变量、JVM 系统属性、database.properties 默认值。database.properties 中不保存真实密码，默认密码字段为 PLEASE_SET_ENV，当未提供真实配置时系统会快速报错，避免使用错误配置继续运行。")

    add_heading(doc, "第四章 业务层设计", 1)
    add_heading(doc, "4.1 业务层职责", 2)
    add_paragraph(doc, "业务层位于 Servlet 和 DAO 之间，主要负责参数标准化、合法性校验、业务状态判断、结果封装和跨 DAO 流程编排。系统使用 ServiceResult 统一表示业务操作是否成功、提示消息和返回数据，使 Servlet 可以用一致方式处理成功与失败响应。")
    add_heading(doc, "4.2 用户与认证业务", 2)
    add_paragraph(doc, "UserService 负责注册和登录业务。注册时先检查用户名、昵称、邮箱和密码格式，再统一转换大小写、校验重复性，最后使用 PasswordUtils 生成 BCrypt 哈希并写入 users 表。登录时根据用户名查询用户，校验密码哈希和账号状态。RememberMeService 负责生成、校验和清理持久化登录令牌，令牌表只保存 selector 和 token_hash，不保存明文密码。")
    add_heading(doc, "4.3 校园广场业务", 2)
    add_paragraph(doc, "PostService 和 SquareService 负责校园广场相关业务，包括帖子列表、发布、详情、评论、点赞、收藏、浏览量更新和内容删除。系统对标题、正文、分类、图片字段进行校验，并在点赞、收藏、评论时维护统计字段。举报业务由 ReportService 负责，举报创建后可通知管理员在后台处理。")
    add_heading(doc, "4.4 二手市场与模拟支付业务", 2)
    add_paragraph(doc, "GoodsService 负责商品发布、编辑、上下架、分类筛选、收藏和状态管理。TradeOrderService 与 MockPayService 负责模拟线上支付流程：买家创建订单后生成支付令牌和二维码链接，确认支付时检查订单状态、过期时间、商品是否仍在售，成功后将订单置为 paid，并将商品状态置为 sold。该流程虽然不接入真实支付平台，但完整模拟了订单状态流转和并发校验。")
    add_heading(doc, "4.5 失物招领与活动业务", 2)
    add_paragraph(doc, "LostFoundService 负责失物招领信息发布、编辑、状态变更和详情查询。ClaimRequestService 负责认领申请，申请提交后将记录状态置为 pending，并根据处理结果更新 lost_found 状态为 claiming、completed 或 pending。ActivityService 负责活动发布和状态管理，ActivityRegistrationService 负责报名和取消报名，报名时校验活动状态、截止时间、容量限制和重复报名。")
    add_heading(doc, "4.6 消息、搜索与后台业务", 2)
    add_paragraph(doc, "MessageService 统一处理站内通知，包括评论、点赞、收藏、认领、活动和举报处理相关通知，并提供未读数量和标记已读功能。SearchService 聚合帖子、商品、失物招领、活动和公告结果，形成全站搜索页面。AdminService 负责后台管理，包括平台统计、用户启停、密码重置、内容状态修改、公告发布和举报审核，是管理员端的核心业务入口。")

    add_heading(doc, "第五章 控制层设计", 1)
    add_heading(doc, "5.1 Servlet 控制器设计", 2)
    add_paragraph(doc, "控制层以 Servlet 为核心，web.xml 中显式配置了登录、注册、首页、内容片段、后台管理、帖子、商品、交易、失物招领、活动、私信、公告、举报等路径。每个 Servlet 只负责 HTTP 层工作：获取请求参数、读取当前登录用户、调用对应 Service、处理异常并返回 JSP 或 JSON。这样可以避免业务逻辑直接写在 JSP 中，也便于单元测试业务层。")
    add_heading(doc, "5.2 过滤器设计", 2)
    add_simple_table(doc, ["过滤器", "作用"], [
        ["EncodingFilter", "统一设置请求与响应编码，解决中文参数和页面输出问题"],
        ["RememberMeFilter", "在 Session 缺失时根据 Cookie 中的持久化令牌尝试自动登录"],
        ["AuthFilter", "保护需要登录的学生端功能，如个人中心、发布、收藏和私信"],
        ["AdminAuthFilter", "保护 /admin 和 /admin/* 后台路径，限制非管理员访问"],
    ], [Cm(4), Cm(12)])
    add_caption(doc, "表5.1 过滤器职责说明")
    add_heading(doc, "5.3 JSON 接口与异步交互", 2)
    add_paragraph(doc, "系统前端大量使用 Fetch API 提交异步请求，控制层通过 PostJsonSupport、GoodsJsonSupport、LostFoundJsonSupport、ActivityJsonSupport、PrivateMessageJsonSupport 和 MessageJsonSupport 等辅助类统一输出 JSON。统一的 JSON 响应结构使前端可以用一致方式展示成功提示、失败原因和局部刷新结果。")
    add_heading(doc, "5.4 文件上传控制", 2)
    add_paragraph(doc, "ImageUploadServlet 负责头像、帖子图片、商品图片、失物招领图片和活动封面的上传。Servlet 通过 multipart-config 限制单文件大小为 5MB、请求总大小为 6MB，并在保存前校验文件类型和业务场景。上传成功后返回可访问路径，页面再将图片路径写入对应业务表。")
    add_caption(doc, "图5.1 Servlet 与过滤器配置截图（截图待补）")

    add_heading(doc, "第六章 视图层设计", 1)
    add_heading(doc, "6.1 页面结构设计", 2)
    add_paragraph(doc, "视图层由 JSP、CSS、JavaScript 和图片资源组成。index.jsp 负责主框架布局，包括顶部导航、左侧导航、全局搜索、消息入口和内容容器；WEB-INF/views 下存放登录、注册、首页、详情页、私信页、后台页和各业务片段；fragments 目录用于首页动态内容区域的局部加载。")
    add_heading(doc, "6.2 前端路由与交互", 2)
    add_paragraph(doc, "前端采用 SPA-Lite 方式组织，app-router.js 根据 URL Hash 和导航项向 ContentServlet 请求 JSP 片段，实现无刷新切换。各业务页面配套独立 JavaScript 文件，例如 market.js 处理商品筛选、发布和收藏，lostfound-actions.js 处理认领和状态变更，activity-actions.js 处理活动报名，private-messages.js 处理会话刷新和消息发送。")
    add_heading(doc, "6.3 页面功能展示", 2)
    add_paragraph(doc, "学生端页面突出校园使用场景：首页展示校园地图、帖子信息流、每日签到、公告、活动推荐和失物速递；校园广场展示帖子列表和互动入口；二手市场通过分类、状态、交易方式和排序提高商品检索效率；失物招领和活动模块分别围绕认领流程与报名流程设计；个人中心集中展示资料、收藏、发布内容、购买记录和签到信息。")
    add_paragraph(doc, "后台管理页面采用独立管理视图，提供统计卡片、用户管理、帖子管理、商品管理、失物招领管理、活动管理、公告管理和举报处理。管理员可以通过统一后台完成内容审核、公告维护和异常内容处理。")
    add_heading(doc, "6.4 静态资源管理", 2)
    add_paragraph(doc, "项目将 CSS、JavaScript 和图片资源按类型分别存放在 src/main/webapp/css、src/main/webapp/js 和 src/main/webapp/images 下。当前包含 14 个 CSS 文件和 14 个 JavaScript 文件，各业务模块尽量使用独立样式和脚本，避免所有交互堆叠在单一文件中。页面输出用户内容时使用 HtmlUtils.escape 进行转义，降低 XSS 风险。")
    add_caption(doc, "图6.1 首页主界面截图（截图待补）")
    add_caption(doc, "图6.2 校园广场帖子列表与详情截图（截图待补）")
    add_caption(doc, "图6.3 二手市场商品列表与支付二维码截图（截图待补）")
    add_caption(doc, "图6.4 失物招领认领流程截图（截图待补）")
    add_caption(doc, "图6.5 校园活动报名页面截图（截图待补）")
    add_caption(doc, "图6.6 后台管理页面截图（截图待补）")

    add_heading(doc, "第七章 项目总结", 1)
    add_heading(doc, "7.1 项目完成情况", 2)
    add_paragraph(doc, "本项目已完成校园综合社区平台的主要功能，覆盖学生端和管理员端两个使用场景。学生端实现了注册登录、校园广场、二手交易、失物招领、活动报名、签到等级、站内消息、私信、个人中心和全站搜索；管理员端实现了统计面板、用户管理、内容管理、公告维护和举报审核。项目能够通过 Maven 打包为 CampusHub.war，并部署到 Tomcat 9 运行。")
    add_heading(doc, "7.2 技术收获", 2)
    add_paragraph(doc, "通过本项目开发，进一步熟悉了 Java Web 应用从请求入口到数据库持久化的完整链路，掌握了 Servlet、Filter、JSP、JDBC、Session、Cookie、Maven 和 MySQL 的综合使用方法。项目中多处业务使用事务处理和状态机思维，例如活动报名人数更新、模拟订单支付、失物认领处理和账号注销清理，这些实践提升了对业务一致性和异常回滚的理解。")
    add_heading(doc, "7.3 存在不足", 2)
    add_paragraph(doc, "项目仍存在一些可以继续优化的地方。第一，前端仍以 JSP 和原生 JavaScript 为主，复杂页面中仍有一定服务端脚本片段，后续可以进一步通过 JSTL、EL 或前后端分离方式降低页面复杂度。第二，图片上传目前保存在应用目录下，适合课程演示，但不适合生产环境长期保存。第三，当前模拟支付没有接入真实第三方支付接口，只完成课程项目层面的订单状态演示。第四，自动化测试主要集中在业务层和工具类，完整浏览器端到端测试还可以继续补充。")
    add_heading(doc, "7.4 后续改进方向", 2)
    add_paragraph(doc, "后续可以从四个方向继续完善：一是引入更系统的前端组件化方案，提升页面复用和维护效率；二是将上传文件迁移到独立持久化目录或对象存储，避免应用重部署导致文件丢失；三是补充端到端自动化测试脚本，覆盖注册、发帖、交易、认领、报名和后台审核等关键流程；四是完善日志、监控和异常告警，使系统具备更好的运行可观测性。")

    return doc


if __name__ == "__main__":
    document = build_doc()
    document.save(OUT)
    print(OUT)
